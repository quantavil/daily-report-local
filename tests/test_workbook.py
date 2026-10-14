from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Font

from daily_report.models import RowUpdate
from daily_report.workbook import WorkbookStore

HEADERS = ['Shipping Bill No', 'SB Date', 'Port Code', 'Document Date', 'Inv No', 'LEO Date',
           'Shipper Name', 'Container No', 'Container Size', 'Customs Seal No', 'Billing Status', 'Status']


def book(tmp_path, rows=None, merges=()):
    path = tmp_path / 'report.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'FY 2026-27'
    ws.append(HEADERS)
    for row in rows or [[7546404.0, date(2026, 10, 7), 'INMBD6']]:
        ws.append(row)
    for cell_range in merges:
        ws.merge_cells(cell_range)
    wb.save(path)
    return path


def update(row=2, status='Success', **values):
    return RowUpdate(row_idx=row, values=values, status=status)


def test_reads_identity_dates_blanks_and_cancelled(tmp_path):
    path = book(tmp_path, [[7546404.0, '07/10/2026', 'inmbd6', '08-Oct-2026', ' N/A ', None,
                            None, None, None, None, None],
                           ['7546405.0', datetime(2026, 10, 7), 'INMBD6'],
                           ['SB-7546406', date(2026, 10, 7), 'INMBD6']])
    rows = WorkbookStore(path).rows('FY 2026-27')
    assert len(rows) == 2
    assert rows[0].identity.sb_no == '7546404'
    assert rows[0].identity.sb_date == date(2026, 10, 7)
    assert rows[0].identity.document_date == date(2026, 10, 8)
    assert rows[0].identity.port_code == 'INMBD6'
    assert rows[0].values['inv_no'] is None
    assert not rows[0].cancelled


def test_exact_sheet_and_required_columns(tmp_path):
    path = book(tmp_path)
    store = WorkbookStore(path)
    with pytest.raises(ValueError, match='Unknown sheet'):
        store.rows('fy 2026-27')
    wb = openpyxl.load_workbook(path)
    wb.active.cell(1, 5).value = 'Unrelated'
    wb.save(path)
    with pytest.raises(ValueError, match='inv_no'):
        WorkbookStore(path).rows('FY 2026-27')


def test_preserves_formulas_and_existing_dates_even_with_blank_force(tmp_path):
    path = book(tmp_path, [[7546404, date(2026, 10, 7), 'INMBD6', None, '=1+1', date(2026, 10, 8)]])
    store = WorkbookStore(path)
    store.apply('FY 2026-27', [update(inv_no='ABC', leo_date=None)], force=True)
    store.save()
    ws = openpyxl.load_workbook(path).active
    assert ws['E2'].value == '=1+1'
    assert ws['F2'].value.date() == date(2026, 10, 8)


def test_force_and_field_scope_date_and_multiline_formatting(tmp_path):
    path = book(tmp_path, [[7546404, date(2026, 10, 7), 'INMBD6', None, 'OLD']])
    wb = openpyxl.load_workbook(path)
    wb.active['F2'].font = Font(bold=True)
    wb.save(path)
    store = WorkbookStore(path)
    store.apply('FY 2026-27', [update(inv_no='NEW', leo_date=date(2026, 10, 9), container_no='A\n\nB')],
                force=True, fields={'leo_date', 'container_no'})
    store.save()
    ws = openpyxl.load_workbook(path).active
    assert ws['E2'].value == 'OLD'
    assert ws['F2'].value.date() == date(2026, 10, 9)
    assert ws['F2'].font.bold
    assert ws['H2'].value == 'A\n\nB'
    assert ws['H2'].alignment.wrap_text


@pytest.mark.parametrize('reverse', [False, True])
def test_merged_conflicts_are_preflighted_and_mark_all_rows_partial(tmp_path, reverse):
    path = book(tmp_path, [[7546404], [7546405]], merges=['E2:E3'])
    store = WorkbookStore(path)
    changes = [update(2, inv_no='A'), update(3, inv_no='B')]
    problems = store.apply('FY 2026-27', changes[::-1] if reverse else changes, force=True)
    assert set(problems) == {2, 3}
    assert all(change.status == 'Partial' for change in changes)
    store.save()
    ws = openpyxl.load_workbook(path).active
    assert ws['E2'].value is None
    assert ws['L2'].value == ws['L3'].value == 'Partial'


def test_merged_child_cannot_mutate_unselected_parent_data_or_status(tmp_path):
    path = book(tmp_path, [[7546404, None, None, None, 'PARENT'], [7546405]], merges=['E2:E3', 'L2:L3'])
    store = WorkbookStore(path)
    changes = [update(3, inv_no='NEW')]
    problems = store.apply('FY 2026-27', changes, force=True)
    store.save()
    ws = openpyxl.load_workbook(path).active
    assert ws['E2'].value == 'PARENT'
    assert ws['L2'].value is None
    assert 3 in problems
    assert changes[0].status == 'Partial'


def test_merged_parent_cannot_mutate_unselected_child(tmp_path):
    path = book(tmp_path, [[7546404], [7546405]], merges=['E2:E3'])
    store = WorkbookStore(path)
    assert 2 in store.apply('FY 2026-27', [update(2, inv_no='NEW')], force=True)
    store.save()
    assert openpyxl.load_workbook(path).active['E2'].value is None


def test_cancelled_updates_only_stamp_status(tmp_path):
    path = book(tmp_path, [[7546404, None, None, None, None, None, None, None, None, None, 'CANCELLED']])
    store = WorkbookStore(path)
    store.apply('FY 2026-27', [update(inv_no='ABC')])
    store.save()
    ws = openpyxl.load_workbook(path).active
    assert ws['E2'].value is None
    assert ws['L2'].value == 'Cancelled'


def test_save_only_when_dirty_atomic_backup_and_retention(tmp_path):
    path = book(tmp_path)
    original = path.read_bytes()
    store = WorkbookStore(path)
    store.save()
    assert not list(tmp_path.glob('*.bak-*'))
    for n in range(7):
        store.apply('FY 2026-27', [update(inv_no=str(n))], force=True)
        store.save()
    backups = sorted(tmp_path.glob('*.bak-*.xlsx'))
    assert len(backups) == 5
    assert len(backups[0].stem.rsplit('-', 1)[-1]) == 6
    assert path.read_bytes() != original
    assert not list(tmp_path.glob('*.tmp.xlsx'))


def test_bad_identity_date_is_not_silently_unknown(tmp_path):
    path = book(tmp_path, [[7546404, 'impossible date', 'INMBD6']])
    rows = WorkbookStore(path).rows('FY 2026-27')
    assert rows[0].error is not None
    assert 'row 2' in rows[0].error
    assert rows[0].identity.sb_date is None


def test_atomic_save_failure_keeps_original(tmp_path, monkeypatch):
    path = book(tmp_path)
    original = path.read_bytes()
    store = WorkbookStore(path)
    store.apply('FY 2026-27', [update(inv_no='NEW')])

    def fail_replace(source: Path, target: Path):
        raise OSError('replace failed')

    monkeypatch.setattr('daily_report.workbook.os.replace', fail_replace)
    with pytest.raises(OSError, match='replace failed'):
        store.save()
    assert path.read_bytes() == original


def test_reads_merged_identity_and_effective_values(tmp_path):
    path = book(tmp_path, [[7546404, date(2026, 10, 7), 'INMBD6', None, 'SHARED'], [None]],
                merges=['A2:A3', 'B2:B3', 'C2:C3', 'E2:E3'])
    rows = WorkbookStore(path).rows('FY 2026-27')
    assert [row.row_idx for row in rows] == [2, 3]
    assert rows[0].identity == rows[1].identity
    assert rows[1].values['inv_no'] == 'SHARED'


def test_missing_status_column_only_created_when_applying(tmp_path):
    path = book(tmp_path)
    wb = openpyxl.load_workbook(path)
    wb.active.delete_cols(12)
    wb.save(path)
    store = WorkbookStore(path)
    store.rows('FY 2026-27')
    store.save()
    assert openpyxl.load_workbook(path).active.max_column == 11
    assert not list(tmp_path.glob('*.bak-*'))
    store.apply('FY 2026-27', [update(status='Not found')])
    store.save()
    ws = openpyxl.load_workbook(path).active
    assert ws['L1'].value == 'Status'
    assert ws['L2'].value == 'Not found'


def test_nonforce_preserves_existing_values_and_unselected_status(tmp_path):
    path = book(tmp_path, [[7546404, None, None, None, 'OLD'], [7546405, None, None, None, 'OTHER']])
    store = WorkbookStore(path)
    store.apply('FY 2026-27', [update(inv_no='NEW')])
    store.save()
    ws = openpyxl.load_workbook(path).active
    assert ws['E2'].value == 'OLD'
    assert ws['E3'].value == 'OTHER'
    assert ws['L3'].value is None


def test_shared_status_conflict_preserves_original(tmp_path):
    path = book(tmp_path, [[7546404], [7546405]], merges=['L2:L3'])
    store = WorkbookStore(path)
    changes = [update(2), update(3, status='Not found')]
    problems = store.apply('FY 2026-27', changes)
    assert set(problems) == {2, 3}
    store.save()
    assert openpyxl.load_workbook(path).active['L2'].value is None
    assert all(change.status == 'Partial' for change in changes)


def test_container_records_serialize_at_write_boundary_with_missing_positions(tmp_path):
    from daily_report.models import ContainerRecord

    path = book(tmp_path)
    store = WorkbookStore(path)
    change = update()
    change.container_fields = {'container_no', 'container_size', 'customs_seal'}
    change.containers = [ContainerRecord(number='ABCU1234567', seal='SEAL-A'),
                         ContainerRecord(number='DEFU7654321', size='40')]
    store.apply('FY 2026-27', [change])
    store.save()
    ws = openpyxl.load_workbook(path).active
    assert ws['H2'].value == 'ABCU1234567\nDEFU7654321'
    assert ws['I2'].value == '\n40'
    assert ws['J2'].value == 'SEAL-A\n'


def test_container_records_only_write_validated_selected_columns(tmp_path):
    from daily_report.models import ContainerRecord

    path = book(tmp_path)
    store = WorkbookStore(path)
    change = update()
    change.container_fields = {'container_no', 'container_size'}
    change.containers = [ContainerRecord(number='ABCU1234567', size='40', seal='SEAL-A')]
    store.apply('FY 2026-27', [change], fields={'container_no', 'customs_seal'})
    store.save()
    ws = openpyxl.load_workbook(path).active
    assert ws['H2'].value == 'ABCU1234567'
    assert ws['I2'].value is None
    assert ws['J2'].value is None


def test_blocked_merged_container_protects_entire_record_triplet(tmp_path):
    from daily_report.models import ContainerRecord
    path = book(tmp_path, [[7546404], [7546404]], merges=['H2:H3'])
    store = WorkbookStore(path)
    change = update()
    change.containers = [ContainerRecord(number='ABCU1234567', size='20', seal='SEAL1')]
    change.container_fields = {'container_no', 'container_size', 'customs_seal'}
    store.apply('FY 2026-27', [change])
    store.save()
    ws = openpyxl.load_workbook(path).active
    assert [ws[cell].value for cell in ('H2', 'I2', 'J2')] == [None, None, None]
    assert change.status == 'Partial'


def test_missing_container_records_rejected_before_any_write(tmp_path):
    path = book(tmp_path)
    store = WorkbookStore(path)
    change = update(inv_no='NEW')
    change.container_fields = {'container_no'}
    with pytest.raises(ValueError, match='Container updates require ordered records'):
        store.apply('FY 2026-27', [change])
    store.save()
    ws = openpyxl.load_workbook(path).active
    assert ws['H2'].value is None
    assert ws['E2'].value is None
    assert ws['L2'].value is None
    assert not list(tmp_path.glob('*.bak-*'))


def test_cancelled_rows_do_not_validate_dates_port_or_existing_data(tmp_path):
    path = book(tmp_path, [[7546404, 'CANCELLED', 'BAD PORT', 'BAD DATE', None, 'BAD LEO',
                            None, None, None, None, 'CANCELLED']])
    rows = WorkbookStore(path).rows('FY 2026-27')
    assert len(rows) == 1
    assert rows[0].cancelled
    assert rows[0].identity.sb_no == '7546404'
    assert rows[0].identity.sb_date is None
    assert rows[0].values == {}


@pytest.mark.parametrize('scope', [{'row_numbers': {3}}, {'sb_no': '7546405'},
                                   {'row_numbers': {2, 3}, 'sb_no': '7546405'}])
def test_unselected_malformed_identity_does_not_block_scoped_reads(tmp_path, scope):
    path = book(tmp_path, [[7546404, 'INVALID DATE', 'BAD PORT'],
                           [7546405, date(2026, 10, 7), 'INMBD6']])
    rows = WorkbookStore(path).rows('FY 2026-27', **scope)
    assert [row.row_idx for row in rows] == [3]
    assert rows[0].identity.sb_no == '7546405'


@pytest.mark.parametrize('cancelled_row', [2, 3])
def test_merged_data_spanning_selected_cancelled_row_is_preserved(tmp_path, cancelled_row):
    rows = [[7546404, None, None, None, 'ORIGINAL'], [7546405]]
    while len(rows[cancelled_row - 2]) < 11:
        rows[cancelled_row - 2].append(None)
    rows[cancelled_row - 2][10] = 'CANCELLED'
    path = book(tmp_path, rows, merges=['E2:E3'])
    store = WorkbookStore(path)
    active_row = 3 if cancelled_row == 2 else 2
    changes = [update(cancelled_row, inv_no='NEVER'), update(active_row, inv_no='NEW')]
    problems = store.apply('FY 2026-27', changes, force=True)
    assert active_row in problems
    assert changes[0].status == 'Cancelled'
    assert changes[1].status == 'Partial'
    store.save()
    ws = openpyxl.load_workbook(path).active
    assert ws['E2'].value == 'ORIGINAL'
    assert ws.cell(cancelled_row, 12).value == 'Cancelled'


def test_invalid_existing_leo_is_retained_without_inferring_cancellation(tmp_path):
    path = book(tmp_path, [[7546404, date(2026, 10, 7), 'INMBD6', None, None, 'CANCELLED',
                            None, None, None, None, 3]])
    original = path.read_bytes()
    store = WorkbookStore(path)
    rows = store.rows('FY 2026-27')
    assert rows[0].values['leo_date'] == 'CANCELLED'
    assert rows[0].identity.sb_date == date(2026, 10, 7)
    assert not rows[0].cancelled
    store.save()
    assert path.read_bytes() == original
    assert not list(tmp_path.glob('*.bak-*'))


def test_force_leo_only_repairs_invalid_existing_leo(tmp_path):
    path = book(tmp_path, [[7546404, date(2026, 10, 7), 'INMBD6', None, 'OLD', 'INVALID DATE']])
    store = WorkbookStore(path)
    assert store.rows('FY 2026-27')[0].values['leo_date'] == 'INVALID DATE'
    store.apply('FY 2026-27', [update(leo_date=date(2026, 10, 9), inv_no='NEW')],
                force=True, fields={'leo_date'})
    store.save()
    ws = openpyxl.load_workbook(path).active
    assert ws['F2'].value.date() == date(2026, 10, 9)
    assert ws['E2'].value == 'OLD'


def test_invalid_identity_row_error_does_not_block_other_selected_rows(tmp_path):
    path = book(tmp_path, [[7546404, 'INVALID DATE', 'INMBD6'],
                           [7546405, date(2026, 10, 7), 'INMBD6']])
    rows = WorkbookStore(path).rows('FY 2026-27')
    assert len(rows) == 2
    assert rows[0].error is not None
    assert rows[0].values == {}
    assert rows[1].error is None
    assert rows[1].identity.sb_date == date(2026, 10, 7)


@pytest.mark.parametrize('status', ['Success', 'Partial', 'Cancelled'])
def test_reads_workbook_status_separately_without_inferring_cancellation(tmp_path, status):
    path = book(tmp_path, [[7546404, date(2026, 10, 7), 'INMBD6', None, None, None,
                            None, None, None, None, 'ACTIVE', status]])
    original = path.read_bytes()
    store = WorkbookStore(path)
    row = store.rows('FY 2026-27')[0]
    assert row.status == status
    assert 'status' not in row.values
    assert not row.cancelled
    store.save()
    assert path.read_bytes() == original
    assert not list(tmp_path.glob('*.bak-*'))


def test_reads_existing_status_for_cancelled_and_identity_error_rows(tmp_path):
    path = book(tmp_path, [[7546404, 'INVALID', 'BAD PORT', None, None, None,
                            None, None, None, None, 'CANCELLED', 'Partial'],
                           [7546405, 'INVALID', 'INMBD6', None, None, None,
                            None, None, None, None, 'ACTIVE', 'Success']])
    rows = WorkbookStore(path).rows('FY 2026-27')
    assert rows[0].status == 'Partial'
    assert rows[0].cancelled
    assert rows[1].status == 'Success'
    assert rows[1].error is not None
    assert not rows[1].cancelled
