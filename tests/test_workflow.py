from datetime import date

import pytest

from daily_report.models import (
    Attachment,
    BillIdentity,
    ContainerRecord,
    DocumentLayout,
    Evidence,
    FieldResult,
    LayoutPage,
    LayoutSpan,
    ParsedDocument,
    WorkbookRow,
)
from daily_report.workflow import make_update, process_bill

IDENTITY = BillIdentity(sb_no='7546404', sb_date=date(2026, 10, 7), port_code='INMBD6')


def attachment(code='SB22', stamp='081020261717', port='INMBD6', message='m'):
    filename = f'754640407102026{port}{code}{stamp}.pdf'
    from daily_report.rules import parse_filename
    meta = parse_filename(filename)
    return Attachment(metadata=meta, message_id=message, attachment_id='a')


class Mail:
    def __init__(self, attachments):
        self.attachments = attachments
        self.downloads = 0
        self.discoveries = 0

    def discover(self, identity, codes):
        self.discoveries += 1
        return [a for a in self.attachments if a.metadata.code in codes]

    def download(self, item):
        self.downloads += 1
        return b'%PDF-' + item.metadata.filename.encode() + item.message_id.encode()


@pytest.fixture
def parser(monkeypatch):
    from daily_report import workflow
    calls = {'layout': 0, 'parse': 0}

    def layout(raw, meta):
        calls['layout'] += 1
        return DocumentLayout(pages=[LayoutPage(number=1, width=600, height=800,
             spans=[LayoutSpan(text='Evidence', bbox=(10, 10, 100, 20))])])

    def parse(layout, meta, digest):
        calls['parse'] += 1
        evidence = Evidence(filename=meta.filename, content_hash=digest, page_number=1,
                            label='Source', quote='Evidence', bbox=(10, 10, 100, 20))
        values = {'SB21': {'shipper_name': 'Exporter', 'inv_no': 'INV/1'},
                  'SB22': {'leo_date': date(2026, 10, 8)},
                  'SB35': {'container_no': 'ABCU1234567', 'container_size': '20', 'customs_seal': 'SEAL1'}}
        fields = {name: FieldResult(state='accepted', value_type='records', evidence=[evidence]) if meta.code == 'SB35' else FieldResult(state='accepted', value=value, evidence=[evidence])
                  for name, value in values[meta.code].items()}
        containers = [ContainerRecord(number='ABCU1234567', size='20', seal='SEAL1', evidence=[evidence])] if meta.code == 'SB35' else []
        return ParsedDocument(metadata=meta, content_hash=digest, fields=fields, containers=containers)

    monkeypatch.setattr(workflow, 'extract_layout', layout)
    monkeypatch.setattr(workflow, 'parse_document', parse)
    monkeypatch.setattr(workflow, 'render_markdown', lambda raw, meta: 'Evidence')
    return calls


def test_first_copy_fills_only_owned_fields(tmp_path, parser):
    result = process_bill(Mail([attachment('SB21')]), tmp_path, IDENTITY)
    assert result.fields['inv_no'].value == 'INV/1'
    assert result.fields['leo_date'].state == 'missing'
    row = WorkbookRow(row_idx=2, identity=IDENTITY)
    update = make_update(row, result)
    assert update.status == 'Partial'
    assert set(update.values) == {'inv_no', 'shipper_name'}


def test_incremental_reuses_artifacts_but_discovers_fresh(tmp_path, parser):
    mail = Mail([attachment('SB22')])
    first = process_bill(mail, tmp_path, IDENTITY, {'leo_date'})
    second = process_bill(mail, tmp_path, IDENTITY, {'leo_date'})
    assert first == second
    assert mail.discoveries == 2
    assert mail.downloads == 1
    assert parser == {'layout': 1, 'parse': 2}


def test_refresh_reparses_without_redownloading(tmp_path, parser):
    mail = Mail([attachment('SB22')])
    process_bill(mail, tmp_path, IDENTITY, {'leo_date'})
    process_bill(mail, tmp_path, IDENTITY, {'leo_date'}, refresh=True)
    assert mail.downloads == 1
    assert parser['parse'] == 2


def test_filename_timestamp_beats_mail_arrival_order(tmp_path, parser):
    older = attachment(stamp='071020261000', message='late-email')
    newer = attachment(stamp='081020261000', message='early-email')
    mail = Mail([older, newer])
    result = process_bill(mail, tmp_path, IDENTITY, {'leo_date'})
    assert result.selected_sources[0].filename == newer.metadata.filename
    assert len(list(tmp_path.rglob('*.pdf'))) == 1
    assert mail.downloads == 1


def test_unexpected_filename_content_conflict_is_explicit_error(tmp_path, parser):
    mail = Mail([attachment(message='m1'), attachment(message='m2')])
    result = process_bill(mail, tmp_path, IDENTITY, {'leo_date'})
    assert result.fields['leo_date'].state == 'error'
    assert parser['parse'] == 0


def test_equal_timestamp_identical_content_is_duplicate(tmp_path, parser):
    mail = Mail([attachment(message='m1'), attachment(message='m2')])
    mail.download = lambda item: b'identical'
    result = process_bill(mail, tmp_path, IDENTITY, {'leo_date'})
    assert result.fields['leo_date'].state == 'accepted'


def test_wrong_port_never_processed(tmp_path, parser):
    result = process_bill(Mail([attachment(port='INXXX6')]), tmp_path, IDENTITY)
    assert all(field.state == 'missing' for field in result.fields.values())
    assert parser['parse'] == 0


def test_unknown_port_multiple_bill_identities_reject(tmp_path, parser):
    expected = IDENTITY.model_copy(update={'port_code': None})
    result = process_bill(Mail([attachment(), attachment(port='INXXX6')]), tmp_path, expected)
    assert all(field.state == 'rejected' for field in result.fields.values())
    assert parser['parse'] == 0


def test_incomplete_workbook_identity_uses_resolved_filename_batch(tmp_path, parser):
    import json
    expected = IDENTITY.model_copy(update={'port_code': None})
    result = process_bill(Mail([attachment()]), tmp_path, expected, {'leo_date'})
    assert result.fields['leo_date'].state == 'accepted'
    payload = json.loads((tmp_path / '2026' / '754640407102026INMBD6' / 'result.json').read_text())
    assert payload['identity']['port_code'] is None
    assert payload['selected_sources'][0]['identity']['port_code'] == 'INMBD6'
    assert {path.name for path in (tmp_path / '2026').iterdir()} == {'754640407102026INMBD6'}


def test_conversion_failure_preserves_existing_date(tmp_path, parser, monkeypatch):
    from daily_report import workflow
    monkeypatch.setattr(workflow, 'extract_layout', lambda *a: (_ for _ in ()).throw(ValueError('OCR disabled')))
    result = process_bill(Mail([attachment()]), tmp_path, IDENTITY, {'leo_date'})
    update = make_update(WorkbookRow(row_idx=2, identity=IDENTITY, values={'leo_date': date(2026, 10, 1)}),
                         result, {'leo_date'}, force=True)
    assert update.status == 'Error'
    assert not update.values


def test_outside_requested_fields_are_never_written(tmp_path, parser):
    result = process_bill(Mail([attachment('SB21'), attachment('SB22')]), tmp_path, IDENTITY)
    update = make_update(WorkbookRow(row_idx=2, identity=IDENTITY), result, {'leo_date'}, force=True)
    assert set(update.values) == {'leo_date'}


def test_complete_existing_row_reports_error_on_source_conflict(tmp_path, parser):
    result = process_bill(Mail([attachment(message='m1'), attachment(message='m2')]), tmp_path, IDENTITY, {'leo_date'})
    row = WorkbookRow(row_idx=2, identity=IDENTITY, values={
        'inv_no': 'I', 'shipper_name': 'E', 'leo_date': date(2026, 10, 1),
        'container_no': 'C', 'container_size': '20', 'customs_seal': 'S'})
    assert make_update(row, result, {'leo_date'}, force=True).status == 'Error'


def test_no_documents_not_found(tmp_path, parser):
    result = process_bill(Mail([]), tmp_path, IDENTITY)
    assert make_update(WorkbookRow(row_idx=2, identity=IDENTITY), result).status == 'Not found'


def book(tmp_path, cancelled=False):
    import openpyxl

    from daily_report.workbook import WorkbookStore
    path = tmp_path / 'report.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'FY 2026-27'
    ws.append(['Shipping Bill No', 'Shipping Bill Date', 'Port Code', 'Shipper Name', 'Inv No',
               'LEO Date', 'Container No', 'Container Size', 'Customs Seal', 'Billing Status', 'Status'])
    for _ in range(2):
        ws.append(['7546404', date(2026, 10, 7), 'INMBD6', None, None, None, None, None, None,
                   'CANCELLED' if cancelled else None, None])
    wb.save(path)
    return WorkbookStore(path)


def test_execute_writes_only_selected_rows(tmp_path, parser):
    from daily_report.workflow import execute
    wb = book(tmp_path)
    mail = Mail([attachment('SB21'), attachment('SB22'), attachment('SB35')])
    records = execute(wb, 'FY 2026-27', mail, tmp_path / 'data', selected=wb.rows('FY 2026-27', row_numbers={3}))
    assert [entry[0].row_idx for entry in records] == [3]
    assert records[0][2].status == 'Success'
    rows = wb.rows('FY 2026-27')
    assert rows[0].values['leo_date'] is None
    assert rows[1].values['leo_date'] == date(2026, 10, 8)
    assert rows[1].values['container_no'] == 'ABCU1234567'
    assert rows[1].values['container_size'] == '20'
    assert rows[1].values['customs_seal'] == 'SEAL1'
    assert wb.wb['FY 2026-27'].cell(2, 11).value is None


def test_execute_verification_reparses_without_workbook_save(tmp_path, parser):
    from daily_report.workflow import execute
    wb = book(tmp_path)
    before = wb.path.read_bytes()
    mail = Mail([attachment('SB22')])
    execute(wb, 'FY 2026-27', mail, tmp_path / 'data', selected=wb.rows('FY 2026-27', row_numbers={2}), leo_only=True)
    before = wb.path.read_bytes()
    records = execute(wb, 'FY 2026-27', mail, tmp_path / 'data', selected=wb.rows('FY 2026-27', row_numbers={2}), leo_only=True, verify=True)
    assert records[0][1].fields['leo_date'].value == date(2026, 10, 8)
    assert wb.path.read_bytes() == before
    assert parser['parse'] == 2
    assert len(list(tmp_path.glob('*.bak-*.xlsx'))) == 1


def test_cancelled_rows_do_not_access_gmail(tmp_path, parser):
    from daily_report.workflow import execute
    wb = book(tmp_path, cancelled=True)
    mail = Mail([])
    records = execute(wb, 'FY 2026-27', mail, tmp_path / 'data', selected=wb.rows('FY 2026-27', row_numbers={2}))
    assert records[0][2].status == 'Cancelled'
    assert mail.discoveries == 0
    assert wb.wb['FY 2026-27'].cell(3, 11).value is None


def test_duplicate_workbook_bill_processed_once(tmp_path, parser):
    from daily_report.workflow import execute
    wb = book(tmp_path)
    mail = Mail([attachment('SB22')])
    execute(wb, 'FY 2026-27', mail, tmp_path / 'data', selected=wb.rows('FY 2026-27'), leo_only=True)
    assert mail.discoveries == 1
    assert parser['parse'] == 1


def test_force_leo_only_preserves_other_fields(tmp_path, parser):
    from daily_report.workflow import execute
    wb = book(tmp_path)
    ws = wb.wb['FY 2026-27']
    ws.cell(2, 4).value = 'Existing exporter'
    ws.cell(2, 6).value = date(2026, 10, 1)
    mail = Mail([attachment('SB21'), attachment('SB22')])
    records = execute(wb, 'FY 2026-27', mail, tmp_path / 'data', selected=wb.rows('FY 2026-27', row_numbers={2}), leo_only=True, force=True)
    assert ws.cell(2, 4).value == 'Existing exporter'
    assert ws.cell(2, 6).value == date(2026, 10, 8)
    assert [meta.code for meta in records[0][1].selected_sources] == ['SB22']


def test_discovery_failure_is_error_not_not_found(tmp_path, parser):
    from daily_report.workflow import execute
    wb = book(tmp_path)
    mail = Mail([])
    mail.discover = lambda *args: (_ for _ in ()).throw(RuntimeError('Gmail unavailable'))
    records = execute(wb, 'FY 2026-27', mail, tmp_path / 'data', selected=wb.rows('FY 2026-27', row_numbers={2}))
    assert records[0][2].status == 'Error'
    assert wb.wb['FY 2026-27'].cell(2, 6).value is None


def test_old_download_failure_does_not_block_newest_evidence(tmp_path, parser):
    older = attachment(stamp='071020261000', message='old')
    newer = attachment(stamp='081020261000', message='new')
    mail = Mail([older, newer])
    def download(item):
        if item.message_id == 'old':
            raise RuntimeError('Old attachment unavailable')
        return b'newest'
    mail.download = download
    result = process_bill(mail, tmp_path, IDENTITY, {'leo_date'})
    assert result.fields['leo_date'].state == 'accepted'
    assert result.selected_sources == [newer.metadata]
    update = make_update(WorkbookRow(row_idx=2, identity=IDENTITY), result)
    assert update.values['leo_date'] == date(2026, 10, 8)
    assert update.status == 'Partial'  # Superseded attachments are no longer downloaded.


def test_failed_newest_does_not_fall_back_to_older(tmp_path, parser):
    mail = Mail([attachment(stamp='071020261000', message='old'), attachment(message='new')])
    def download(item):
        if item.message_id == 'new':
            raise RuntimeError('Newest unavailable')
        return b'old'
    mail.download = download
    result = process_bill(mail, tmp_path, IDENTITY, {'leo_date'})
    assert result.fields['leo_date'].state == 'error'
    assert parser['parse'] == 0


def test_missing_seal_never_attaches_to_different_existing_container(tmp_path, parser):
    result = process_bill(Mail([attachment('SB35')]), tmp_path, IDENTITY)
    row = WorkbookRow(row_idx=2, identity=IDENTITY,
                      values={'container_no': 'DEFU7654321', 'container_size': '40'})
    update = make_update(row, result)
    assert not update.container_fields
    assert 'Container records unchanged' in update.detail
    assert update.status == 'Partial'


def test_matching_existing_container_allows_its_missing_seal(tmp_path, parser):
    result = process_bill(Mail([attachment('SB35')]), tmp_path, IDENTITY)
    row = WorkbookRow(row_idx=2, identity=IDENTITY,
                      values={'container_no': 'ABCU1234567', 'container_size': '20'})
    update = make_update(row, result)
    assert update.container_fields == {'customs_seal'}
    assert update.containers[0].seal == 'SEAL1'
    assert not update.values


def test_partial_record_sequence_remains_partial():
    from daily_report.models import BillResult
    ev = Evidence(filename='source.pdf', content_hash='hash', page_number=1,
                  label='12.Seal', quote='X', bbox=(0, 0, 10, 10))
    result = BillResult(identity=IDENTITY, selected_sources=[attachment('SB35').metadata], fields={
        'customs_seal': FieldResult(state='accepted', value_type='records', evidence=[ev])},
        containers=[ContainerRecord(number='ABCU1234567', size='20', evidence=[ev]),
                    ContainerRecord(number='DEFU7654321', size='20', seal='X', evidence=[ev])])
    row = WorkbookRow(row_idx=2, identity=IDENTITY, values={
        'inv_no':'I', 'shipper_name':'E', 'leo_date':date(2026,10,8),
        'container_no':'ABCU1234567\nDEFU7654321', 'container_size':'20\n20'})
    assert make_update(row, result, {'customs_seal'}, force=True).status == 'Partial'


def test_selected_identity_error_does_not_fetch_or_block_other_bill(tmp_path, parser):
    from daily_report.workflow import execute
    wb = book(tmp_path)
    ws = wb.wb['FY 2026-27']
    ws.cell(2, 2).value = 'not a date'
    mail = Mail([attachment('SB22')])
    records = execute(wb, 'FY 2026-27', mail, tmp_path / 'data', selected=wb.rows('FY 2026-27'), leo_only=True)
    assert records[0][2].status == 'Error'
    assert records[1][1].fields['leo_date'].state == 'accepted'
    assert mail.discoveries == 1


def test_force_does_not_pair_new_seal_with_protected_container_formula(tmp_path, parser):
    result = process_bill(Mail([attachment('SB35')]), tmp_path, IDENTITY)
    row = WorkbookRow(row_idx=2, identity=IDENTITY, values={'container_no': '=A2', 'container_size': '20'})
    update = make_update(row, result, force=True)
    assert not update.container_fields
    assert update.status == 'Partial'


def test_execute_uses_already_selected_rows(tmp_path, parser, monkeypatch):
    from daily_report.workflow import execute
    wb = book(tmp_path)
    selected = wb.rows('FY 2026-27', row_numbers={3})
    monkeypatch.setattr(wb, 'rows', lambda *a, **kw: (_ for _ in ()).throw(AssertionError('Selection repeated')))
    records = execute(wb, 'FY 2026-27', Mail([attachment('SB22')]), tmp_path / 'data',
                      selected=selected, leo_only=True)
    assert [row.row_idx for row, _, _ in records] == [3]
    assert records[0][1].fields['leo_date'].value == date(2026, 10, 8)


def test_old_verified_success_reuses_local_without_gmail(tmp_path, parser):
    mail = Mail([attachment('SB21'), attachment('SB22'), attachment('SB35')])
    process_bill(mail, tmp_path, IDENTITY, today=date(2026, 10, 14))
    mail.discover = lambda *args: (_ for _ in ()).throw(AssertionError('Gmail used for old Success'))
    mail.download = lambda *args: (_ for _ in ()).throw(AssertionError('Download used for old Success'))
    result = process_bill(mail, tmp_path, IDENTITY, today=date(2026, 11, 6), row_successful=True)
    assert all(field.state == 'accepted' for field in result.fields.values())
    assert parser == {'layout': 3, 'parse': 6}


def test_old_missing_documents_are_searched_not_frozen(tmp_path, parser):
    mail = Mail([attachment('SB21'), attachment('SB22'), attachment('SB35')])
    result = process_bill(mail, tmp_path, IDENTITY, today=date(2026, 11, 6))
    assert mail.discoveries == 1 and mail.downloads == 3
    assert all(field.state == 'accepted' for field in result.fields.values())


def test_old_partial_result_continues_gmail_to_find_missing_types(tmp_path, parser):
    mail = Mail([attachment('SB21')])
    process_bill(mail, tmp_path, IDENTITY, today=date(2026, 10, 14))
    mail.attachments += [attachment('SB22'), attachment('SB35')]
    result = process_bill(mail, tmp_path, IDENTITY, today=date(2026, 11, 6))
    assert mail.discoveries == 2 and mail.downloads == 3
    assert all(field.state == 'accepted' for field in result.fields.values())


def test_young_verified_success_still_checks_updated_mail(tmp_path, parser):
    mail = Mail([attachment('SB21'), attachment('SB22'), attachment('SB35')])
    process_bill(mail, tmp_path, IDENTITY, today=date(2026, 10, 14))
    process_bill(mail, tmp_path, IDENTITY, today=date(2026, 11, 5))
    assert mail.discoveries == 2 and mail.downloads == 3


def test_old_success_with_corrupt_pdf_checks_gmail_for_repair(tmp_path, parser):
    from daily_report.rules import batch_name
    mail = Mail([attachment('SB21'), attachment('SB22'), attachment('SB35')])
    process_bill(mail, tmp_path, IDENTITY, today=date(2026, 10, 14))
    (tmp_path / '2026' / batch_name(IDENTITY) / 'pdfs' / attachment('SB21').metadata.filename).write_bytes(b'corrupt')
    result = process_bill(mail, tmp_path, IDENTITY, today=date(2026, 11, 6))
    assert mail.discoveries == 2 and mail.downloads == 4
    assert all(field.state == 'accepted' for field in result.fields.values())


def test_new_leo_and_gatepass_replace_prior_copies_without_redownloading_first_copy(tmp_path, parser):
    from daily_report.rules import batch_name
    original = [attachment('SB21'), attachment('SB22'), attachment('SB35')]
    mail = Mail(original)
    process_bill(mail, tmp_path, IDENTITY, today=date(2026, 10, 14))
    mail.attachments = [original[0], attachment('SB22', stamp='091020261700', message='new-leo'),
                        attachment('SB35', stamp='091020261730', message='new-gatepass')]
    result = process_bill(mail, tmp_path, IDENTITY, today=date(2026, 10, 14))
    assert mail.downloads == 5
    assert all(field.state == 'accepted' for field in result.fields.values())
    pdfs = sorted((tmp_path / '2026' / batch_name(IDENTITY) / 'pdfs').glob('*.pdf'))
    assert len(pdfs) == 3
    assert {p.name for p in pdfs} == {a.metadata.filename for a in mail.attachments}
    for old in original[1:]:
        stem = old.metadata.filename[:-4]
        assert not (pdfs[0].parent.parent / 'processed' / (stem + '.json')).exists()
        assert not (pdfs[0].parent.parent / 'processed' / (stem + '.md')).exists()


def test_failed_newest_download_preserves_prior_pdf_without_accepting_old_fields(tmp_path, parser):
    from daily_report.rules import batch_name
    old = attachment(stamp='071020261000', message='old')
    mail = Mail([old])
    process_bill(mail, tmp_path, IDENTITY, {'leo_date'}, today=date(2026, 10, 14))
    mail.attachments = [old, attachment(message='new')]
    mail.download = lambda item: (_ for _ in ()).throw(RuntimeError('Newest download failed'))
    result = process_bill(mail, tmp_path, IDENTITY, {'leo_date'}, today=date(2026, 10, 14))
    assert result.fields['leo_date'].state == 'error'
    assert parser['parse'] == 1
    pdfs = list((tmp_path / '2026' / batch_name(IDENTITY) / 'pdfs').glob('*.pdf'))
    assert [p.name for p in pdfs] == [old.metadata.filename]


def test_old_failed_newest_revision_retries_gmail_without_falling_back(tmp_path, parser):
    old = attachment(stamp='071020261000', message='old')
    newer = attachment(stamp='081020261000', message='new')
    mail = Mail([old])
    process_bill(mail, tmp_path, IDENTITY, {'leo_date'}, today=date(2026, 10, 14))
    mail.attachments = [old, newer]
    mail.download = lambda item: (_ for _ in ()).throw(RuntimeError('Newest download failed'))
    failed = process_bill(mail, tmp_path, IDENTITY, {'leo_date'}, today=date(2026, 11, 5))
    assert failed.fields['leo_date'].state == 'error'
    retried = process_bill(mail, tmp_path, IDENTITY, {'leo_date'}, today=date(2026, 11, 6))
    assert retried.fields['leo_date'].state == 'error'
    assert retried.fields['leo_date'].value is None
    assert mail.discoveries == 3
    assert parser['parse'] == 1


def test_old_corrupt_first_copy_is_repaired_without_blocking_healthy_types(tmp_path, parser):
    from daily_report.rules import batch_name
    mail = Mail([attachment('SB21'), attachment('SB22'), attachment('SB35')])
    process_bill(mail, tmp_path, IDENTITY, today=date(2026, 10, 14))
    (tmp_path / '2026' / batch_name(IDENTITY) / 'pdfs' / attachment('SB21').metadata.filename).write_bytes(b'corrupt')
    result = process_bill(mail, tmp_path, IDENTITY, today=date(2026, 11, 6))
    assert result.fields['shipper_name'].state == 'accepted'
    assert result.fields['inv_no'].state == 'accepted'
    assert result.fields['leo_date'].state == 'accepted'
    assert result.fields['container_no'].state == 'accepted'
    assert mail.discoveries == 2 and mail.downloads == 4


def test_old_excel_partial_with_complete_source_still_checks_gmail_until_success(tmp_path, parser, monkeypatch):
    from datetime import datetime

    from daily_report import workflow
    wb = book(tmp_path)
    mail = Mail([attachment('SB21'), attachment('SB22'), attachment('SB35')])
    workflow.execute(wb, 'FY 2026-27', mail, tmp_path / 'data',
                     selected=wb.rows('FY 2026-27', row_numbers={2}), force=True)
    wb.wb['FY 2026-27'].cell(2, 11).value = 'Partial'
    class Clock:
        @staticmethod
        def now(tz): return datetime(2026, 11, 6, tzinfo=tz)
    monkeypatch.setattr(workflow, 'datetime', Clock)
    workflow.execute(wb, 'FY 2026-27', mail, tmp_path / 'data',
                     selected=wb.rows('FY 2026-27', row_numbers={2}), force=True)
    assert mail.discoveries == 2
    assert wb.wb['FY 2026-27'].cell(2, 11).value == 'Success'
    workflow.execute(wb, 'FY 2026-27', mail, tmp_path / 'data',
                     selected=wb.rows('FY 2026-27', row_numbers={2}), force=True)
    assert mail.discoveries == 2


def test_old_excel_success_without_pdf_proof_still_checks_gmail(tmp_path, parser, monkeypatch):
    from datetime import datetime

    from daily_report import workflow
    wb = book(tmp_path)
    wb.wb['FY 2026-27'].cell(2, 11).value = 'Success'
    class Clock:
        @staticmethod
        def now(tz): return datetime(2026, 11, 6, tzinfo=tz)
    monkeypatch.setattr(workflow, 'datetime', Clock)
    mail = Mail([attachment('SB21'), attachment('SB22'), attachment('SB35')])
    workflow.execute(wb, 'FY 2026-27', mail, tmp_path / 'data',
                     selected=wb.rows('FY 2026-27', row_numbers={2}), force=True)
    assert mail.discoveries == 1 and mail.downloads == 3


def test_known_failed_newest_blocks_old_evidence_even_if_gmail_omits_newer_later(tmp_path, parser):
    old = attachment(stamp='071020261000', message='old')
    newer = attachment(stamp='081020261000', message='new')
    mail = Mail([old])
    process_bill(mail, tmp_path, IDENTITY, {'leo_date'}, today=date(2026, 10, 14))
    mail.attachments = [old, newer]
    mail.download = lambda item: (_ for _ in ()).throw(RuntimeError('Newest download failed'))
    process_bill(mail, tmp_path, IDENTITY, {'leo_date'}, today=date(2026, 11, 5))
    mail.attachments = [old]
    result = process_bill(mail, tmp_path, IDENTITY, {'leo_date'}, today=date(2026, 11, 6))
    assert result.fields['leo_date'].state == 'error'
    assert result.fields['leo_date'].value is None
    assert parser['parse'] == 1


def test_old_success_fresh_parser_rejection_reopens_gmail_without_double_parsing(tmp_path, parser, monkeypatch):
    from daily_report import workflow
    mail = Mail([attachment('SB21'), attachment('SB22'), attachment('SB35')])
    process_bill(mail, tmp_path, IDENTITY, today=date(2026, 10, 14))
    original_parse = workflow.parse_document
    def revised_parse(layout, metadata, digest):
        parsed = original_parse(layout, metadata, digest)
        if metadata.code == 'SB22':
            parsed.fields['leo_date'] = FieldResult(state='rejected', reason='Current layout rule rejects source')
        return parsed
    monkeypatch.setattr(workflow, 'parse_document', revised_parse)
    result = process_bill(mail, tmp_path, IDENTITY, today=date(2026, 11, 6), row_successful=True)
    assert mail.discoveries == 2
    assert result.fields['leo_date'].state == 'rejected'
    assert parser == {'layout': 3, 'parse': 6}


def test_normal_sync_checks_fully_filled_partial_rows(tmp_path, parser):
    from daily_report.workflow import execute
    wb = book(tmp_path)
    mail = Mail([attachment('SB21'), attachment('SB22'), attachment('SB35')])
    execute(wb, 'FY 2026-27', mail, tmp_path / 'data', selected=wb.rows('FY 2026-27', row_numbers={2}), force=True)
    wb.wb['FY 2026-27'].cell(2, 11).value = 'Partial'
    records = execute(wb, 'FY 2026-27', mail, tmp_path / 'data', selected=wb.rows('FY 2026-27', row_numbers={2}))
    assert len(records) == 1 and mail.discoveries == 2
    assert records[0][2].status == 'Success'
    assert not records[0][2].values and not records[0][2].container_fields


def test_filled_excel_mismatch_is_partial_without_force(tmp_path, parser):
    result = process_bill(Mail([attachment('SB21'), attachment('SB22'), attachment('SB35')]), tmp_path, IDENTITY)
    row = WorkbookRow(row_idx=2, identity=IDENTITY, values={
        'shipper_name':'Exporter', 'inv_no':'WRONG', 'leo_date':date(2026,10,8),
        'container_no':'ABCU1234567','container_size':'20','customs_seal':'SEAL1'})
    update = make_update(row, result)
    assert update.status == 'Partial'
    assert 'inv_no' in update.detail
    assert not update.values and not update.container_fields


def test_filled_excel_with_rejected_source_is_partial_without_force(tmp_path, parser):
    result = process_bill(Mail([attachment('SB21'), attachment('SB22'), attachment('SB35')]), tmp_path, IDENTITY)
    result.fields['leo_date'] = FieldResult(state='rejected', reason='Source rejected')
    row = WorkbookRow(row_idx=2, identity=IDENTITY, values={
        'shipper_name':'Exporter', 'inv_no':'INV/1', 'leo_date':date(2026,10,8),
        'container_no':'ABCU1234567','container_size':'20','customs_seal':'SEAL1'})
    update = make_update(row, result)
    assert update.status == 'Partial'
    assert not update.values
