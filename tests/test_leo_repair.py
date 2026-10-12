"""Regression coverage for selecting and repairing only LEO cells."""
from datetime import date, datetime
from unittest.mock import patch

import openpyxl

from excel_manager import ExcelManager
from models import CustomsExtractedFields, LeoDateEvidence
from pdf_processor import PdfConversionError
from sync import cmd_sync

HEADERS = ["Shipping Bill No", "Inv No", "LEO Date", "Shipper Name", "Container No.",
           "Container Size", "Customs Seal Number", "Billing Status", "Status"]
ORIGINAL = ["3924868", "INV-OLD", date(2026, 6, 6), "ORIGINAL SHIPPER", "TCKU7809338",
            "40", "SEAL-OLD", None, None]


def workbook(tmp_path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FY 2026-27"
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    path = tmp_path / "repair.xlsx"
    wb.save(path)
    manager = ExcelManager(path)
    manager.load()
    return manager


class Gmail:
    def __init__(self):
        self.fetched = []

    def authenticate(self):
        return True

    def fetch_documents(self, sb, *, fields=None, include_first_copy=False, first_copy_only=False):
        self.fetched.append(sb)
        return {"finalLeo": ("sb22.pdf", b"external-boundary-placeholder")}


def extracted():
    return CustomsExtractedFields(
        invoice_numbers="INV-NEW", leo_date="08-Jun-2026", shipper_name="NEW SHIPPER",
        container_numbers="UNSU0088111", container_sizes="20", customs_seal_numbers="SEAL-NEW",
    )


def test_force_without_filter_selects_complete_rows(tmp_path, monkeypatch):
    """Regression: --force must include complete rows, not just missing rows."""
    monkeypatch.setenv("OPENROUTER_API_KEY", "fake-unit-test-key")
    manager = workbook(tmp_path, [ORIGINAL])
    gmail = Gmail()
    with patch("sync.process_documents", return_value=extracted()):
        cmd_sync(manager, "FY 2026-27", gmail, force=True, workers=1)
    assert gmail.fetched == ["3924868"]
    loaded = openpyxl.load_workbook(manager.file_path)
    assert loaded.active.cell(2, 3).value == datetime(2026, 6, 8)


def test_leo_only_force_preserves_other_fields(tmp_path, monkeypatch):
    """The AI can return all six fields; only LEO may be overwritten."""
    monkeypatch.setenv("OPENROUTER_API_KEY", "fake-unit-test-key")
    manager = workbook(tmp_path, [ORIGINAL])
    gmail = Gmail()
    with patch("sync.process_documents", return_value=extracted()):
        cmd_sync(manager, "FY 2026-27", gmail, force=True, leo_only=True, workers=1)
    loaded = openpyxl.load_workbook(manager.file_path)
    assert loaded.active.cell(2, 3).value == datetime(2026, 6, 8)
    assert [loaded.active.cell(2, c).value for c in [2, 4, 5, 6, 7]] == [
        "INV-OLD", "ORIGINAL SHIPPER", "TCKU7809338", "40", "SEAL-OLD",
    ]
    assert list(tmp_path.glob("repair.bak-*.xlsx"))


def test_leo_only_does_not_fetch_rows_with_filled_leo(tmp_path, monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "fake-unit-test-key")
    incomplete = ORIGINAL.copy()
    incomplete[1] = None
    manager = workbook(tmp_path, [incomplete])
    gmail = Gmail()
    cmd_sync(manager, "FY 2026-27", gmail, leo_only=True, workers=1)
    assert gmail.fetched == []
    assert manager.wb.active.cell(2, 3).value == datetime(2026, 6, 6)
    assert manager.wb.active.cell(2, 9).value == "Partial"


def test_filtered_repair_does_not_stamp_unselected_cancelled_rows(tmp_path, monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "fake-unit-test-key")
    cancelled = ORIGINAL.copy()
    cancelled[0] = "1234567"
    cancelled[7] = "CANCELLED"
    manager = workbook(tmp_path, [ORIGINAL, cancelled])
    gmail = Gmail()
    with patch("sync.process_documents", return_value=extracted()):
        cmd_sync(manager, "FY 2026-27", gmail, force=True, rows={2}, workers=1)
    assert gmail.fetched == ["3924868"]
    assert manager.wb.active.cell(3, 9).value is None


def test_leo_only_keeps_overall_status_partial_when_other_fields_missing(tmp_path, monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "fake-unit-test-key")
    incomplete = ORIGINAL.copy()
    incomplete[1] = None
    incomplete[2] = None
    manager = workbook(tmp_path, [incomplete])
    gmail = Gmail()
    with patch("sync.process_documents", return_value=extracted()):
        cmd_sync(manager, "FY 2026-27", gmail, leo_only=True, workers=1)
    assert manager.wb.active.cell(2, 3).value == date(2026, 6, 8)
    assert manager.wb.active.cell(2, 2).value is None
    assert manager.wb.active.cell(2, 9).value == "Partial"


def test_forced_leo_missing_evidence_preserves_date_and_flags_partial(tmp_path, monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "fake-unit-test-key")
    manager = workbook(tmp_path, [ORIGINAL])
    with patch("sync.process_documents", return_value=CustomsExtractedFields()):
        cmd_sync(manager, "FY 2026-27", Gmail(), force=True, leo_only=True, workers=1)
    assert manager.wb.active.cell(2, 3).value == datetime(2026, 6, 6)
    assert manager.wb.active.cell(2, 9).value == "Partial"


def test_forced_malformed_ai_keeps_filled_values_and_reports_incomplete(tmp_path, monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "fake-unit-test-key")
    manager = workbook(tmp_path, [ORIGINAL])
    malformed = CustomsExtractedFields(leo_warning="AI returned malformed response")

    with patch("sync.process_documents", return_value=malformed):
        cmd_sync(manager, "FY 2026-27", Gmail(), force=True, workers=1)

    assert manager.wb.active.cell(2, 2).value == "INV-OLD"
    assert manager.wb.active.cell(2, 3).value == datetime(2026, 6, 6)
    assert manager.wb.active.cell(2, 9).value == "Partial"

    from sync import _warning_detail

    assert _warning_detail("AI returned malformed response", ["inv_no"], malformed) == (
        "Extraction incomplete: AI returned malformed response"
    )


def test_malformed_sb21_fallback_keeps_existing_leo_and_partial_status(tmp_path, monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "fake-unit-test-key")
    manager = workbook(tmp_path, [ORIGINAL])

    class FallbackGmail(Gmail):
        def fetch_documents(self, sb, *, fields=None, include_first_copy=False, first_copy_only=False):
            self.fetched.append(sb)
            if include_first_copy or first_copy_only:
                return {"firstCopy": ("sb21.pdf", b"sb21")}
            return {"finalLeo": ("sb22.pdf", b"sb22")}

    def process(_docs, *, fields, **_kwargs):
        if fields == {"inv_no", "shipper_name"}:
            return CustomsExtractedFields(leo_warning="AI returned malformed response")
        return CustomsExtractedFields()

    with patch("sync.process_documents", side_effect=process):
        cmd_sync(manager, "FY 2026-27", FallbackGmail(), force=True, workers=1)

    assert manager.wb.active.cell(2, 3).value == datetime(2026, 6, 6)
    assert manager.wb.active.cell(2, 9).value == "Partial"


def test_conversion_error_with_sb21_fallback_writes_invoice_but_marks_error(
    tmp_path, monkeypatch
):
    monkeypatch.setenv("OPENROUTER_API_KEY", "fake-unit-test-key")
    manager = workbook(tmp_path, [ORIGINAL])

    class ConversionFailureGmail(Gmail):
        def fetch_documents(self, sb, *, fields=None, include_first_copy=False, first_copy_only=False):
            self.fetched.append((sb, include_first_copy, first_copy_only))
            if include_first_copy or first_copy_only:
                return {"firstCopy": ("sb21.pdf", b"sb21")}
            return {
                "eGatepass": ("gatepass.pdf", b"gatepass"),
                "finalLeo": ("sb22.pdf", b"needs-ocr"),
            }

    def trim(raw_bytes, doc_type, *, leo_only=False):
        if raw_bytes == b"needs-ocr":
            raise PdfConversionError("page 2 requires OCR")
        text = "--- SOURCE PAGE 1 ---\nShipping Bill No: 3924868"
        if raw_bytes == b"gatepass":
            text += "\n1.Order No / date 1/101 08-JUN-26"
        else:
            text += "\n2.INV NO. INV-SB21"
        return text, [0]

    def extract(text, doc_type="COMBINED", sb_no="", fields=None):
        if "First Copy Shipping Bill (SB21)" in text:
            return CustomsExtractedFields(invoice_numbers="INV-SB21")
        return CustomsExtractedFields(
            shipper_name="GATEPASS EXPORTER",
            leo_evidence=[{
                "doc_type": "EGATEPASS",
                "page_number": 1,
                "label": "1.Order No / date",
                "raw_date": "08-JUN-26",
                "quote": "1.Order No / date 1/101 08-JUN-26",
            }],
        )

    monkeypatch.setattr("sync.trim_pdf", trim)
    monkeypatch.setattr("sync.extract_with_ai", extract)

    cmd_sync(manager, "FY 2026-27", ConversionFailureGmail(), force=True, workers=1)

    assert manager.wb.active.cell(2, 2).value == "INV-SB21"
    assert manager.wb.active.cell(2, 3).value == datetime(2026, 6, 6)
    assert manager.wb.active.cell(2, 4).value == "GATEPASS EXPORTER"
    assert manager.wb.active.cell(2, 9).value == "Error"


def _assert_primary_data_persisted_after_fallback_error(tmp_path, monkeypatch, failure, capsys):
    monkeypatch.setenv("OPENROUTER_API_KEY", "fake-unit-test-key")
    manager = workbook(tmp_path, [ORIGINAL])
    calls = []

    class FallbackFailureGmail(Gmail):
        def fetch_documents(
            self, sb, *, fields=None, include_first_copy=False, first_copy_only=False
        ):
            calls.append((set(fields or ()), include_first_copy, first_copy_only))
            if first_copy_only:
                if failure == "gmail":
                    raise RuntimeError("fallback Gmail unavailable")
                return {"firstCopy": ("sb21.pdf", b"sb21")}
            return {
                "eGatepass": ("gatepass.pdf", b"gatepass"),
                "finalLeo": ("sb22.pdf", b"sb22"),
            }

    def trim(raw_bytes, _doc_type, *, leo_only=False):
        if raw_bytes == b"gatepass":
            text = (
                "Shipping Bill No: 3924868\n"
                "1.Order No / date 1/101 08-JUN-26\n"
                "4.CHA VALID PRIMARY SHIPPER\nAllowed for Movement"
            )
        elif raw_bytes == b"sb22":
            text = "Shipping Bill No: 3924868\n2.INV NO. INV-SB22"
        else:
            text = "Shipping Bill No: 3924868\n2.INV NO. INV-SB21"
        return f"--- SOURCE PAGE 1 ---\n{text}", [0]

    def extract(text, doc_type="COMBINED", sb_no="", fields=None):
        if "First Copy Shipping Bill (SB21)" in text and failure == "ai":
            raise RuntimeError("fallback OpenRouter unavailable")
        if "First Copy Shipping Bill (SB21)" in text:
            return CustomsExtractedFields(invoice_numbers="INV-SB21")
        return CustomsExtractedFields(
            shipper_name="VALID PRIMARY SHIPPER",
            leo_evidence=[LeoDateEvidence(
                doc_type="EGATEPASS",
                page_number=1,
                label="1.Order No / date",
                raw_date="08-JUN-26",
                quote="1.Order No / date 1/101 08-JUN-26",
            )],
        )

    monkeypatch.setattr("sync.trim_pdf", trim)
    monkeypatch.setattr("sync.extract_with_ai", extract)
    cmd_sync(manager, "FY 2026-27", FallbackFailureGmail(), force=True, workers=1)

    assert calls == [
        ({"inv_no", "leo_date", "shipper_name", "container_no", "container_size", "customs_seal"}, False, False),
        ({"inv_no"}, False, True),
    ]
    saved = openpyxl.load_workbook(manager.file_path).active
    assert saved.cell(2, 2).value == "INV-OLD"
    assert saved.cell(2, 3).value == datetime(2026, 6, 8)
    assert saved.cell(2, 4).value == "VALID PRIMARY SHIPPER"
    assert saved.cell(2, 9).value == "Error"
    output = capsys.readouterr().out
    assert f"fallback {'Gmail' if failure == 'gmail' else 'OpenRouter'} unavailable" in output


def test_fallback_gmail_exception_persists_valid_primary_fields_as_error(
    tmp_path, monkeypatch, capsys
):
    _assert_primary_data_persisted_after_fallback_error(
        tmp_path, monkeypatch, "gmail", capsys
    )


def test_fallback_ai_exception_persists_valid_primary_fields_as_error(
    tmp_path, monkeypatch, capsys
):
    _assert_primary_data_persisted_after_fallback_error(
        tmp_path, monkeypatch, "ai", capsys
    )


def test_invalid_sheet_refuses_fallback_before_writing(tmp_path, monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "fake-unit-test-key")
    import pytest
    import typer

    manager = workbook(tmp_path, [ORIGINAL])
    before = manager.file_path.read_bytes()
    with pytest.raises(typer.BadParameter):
        cmd_sync(manager, "FY TYPO", Gmail(), force=True, workers=1)
    assert manager.file_path.read_bytes() == before
    assert manager.wb.active.cell(2, 9).value is None


def test_leo_audit_prints_evidence_without_saving(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("OPENROUTER_API_KEY", "fake-unit-test-key")
    from models import LeoDateEvidence
    from sync import cmd_verify

    manager = workbook(tmp_path, [ORIGINAL])
    before = manager.file_path.read_bytes()
    fields = CustomsExtractedFields(leo_date="08-Jun-2026", leo_evidence=[LeoDateEvidence(
        doc_type="FINAL_LEO", page_number=1, label="6.LEO Date", raw_date="08-JUN-26",
        quote="6.LEO Date 08-JUN-26",
    )])
    with patch("sync.process_documents", return_value=fields) as process:
        cmd_verify(manager, "FY 2026-27", Gmail(), rows={2}, leo_only=True)
    process.assert_called_once()
    assert process.call_args.kwargs["fields"] == {"leo_date"}
    assert process.call_args.kwargs.get("cache") is None
    assert process.call_args.kwargs["refresh"] is True
    assert manager.file_path.read_bytes() == before
    assert not list(tmp_path.glob("repair.bak-*.xlsx"))
    output = capsys.readouterr().out
    assert "6.LEO Date" in output
    assert "MISMATCH" in output


def test_audit_uses_scoped_sb21_fallback_without_saving(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("OPENROUTER_API_KEY", "fake-unit-test-key")
    from sync import cmd_verify

    manager = workbook(tmp_path, [ORIGINAL])
    before = manager.file_path.read_bytes()
    calls = []

    class FallbackGmail(Gmail):
        def fetch_documents(self, sb, *, fields=None, include_first_copy=False, first_copy_only=False):
            self.fetched.append((sb, include_first_copy, first_copy_only))
            if include_first_copy or first_copy_only:
                return {"firstCopy": ("sb21.pdf", b"sb21")}
            return {"finalLeo": ("sb22.pdf", b"sb22")}

    def process(docs, *, fields, cache=None, refresh=False, **_kwargs):
        calls.append((set(fields), set(docs), cache, refresh))
        if "firstCopy" in docs:
            return CustomsExtractedFields(invoice_numbers="INV-SB21")
        return CustomsExtractedFields(shipper_name="NEW SHIPPER")

    with patch("sync.process_documents", side_effect=process):
        gmail = FallbackGmail()
        cmd_verify(manager, "FY 2026-27", gmail, rows={2})

    output = capsys.readouterr().out
    assert gmail.fetched == [("3924868", False, False), ("3924868", False, True)]
    assert calls[0][1] == {"finalLeo"}
    assert calls[1][0] == {"inv_no"}
    assert calls[1][1] == {"firstCopy"}
    assert calls[0][2:] == (None, True)
    assert calls[1][2:] == (None, True)
    assert "INV-SB21" in output
    assert "MISMATCH" in output
    assert manager.file_path.read_bytes() == before
    assert not list(tmp_path.glob("repair.bak-*.xlsx"))


def test_leo_audit_reports_no_eligible_documents_for_selected_fields(
    tmp_path, monkeypatch, capsys
):
    monkeypatch.setenv("OPENROUTER_API_KEY", "fake-unit-test-key")
    from sync import cmd_verify

    manager = workbook(tmp_path, [ORIGINAL])

    class FirstCopyOnlyGmail(Gmail):
        def fetch_documents(self, sb, *, fields=None, include_first_copy=False, first_copy_only=False):
            return {"firstCopy": ("sb21.pdf", b"sb21")}

    cmd_verify(
        manager,
        "FY 2026-27",
        FirstCopyOnlyGmail(),
        rows={2},
        leo_only=True,
    )

    output = capsys.readouterr().out
    assert "No eligible ICEGATE documents for selected fields" in output


def test_skipping_filled_cells_does_not_clear_a_previous_rejection(tmp_path, monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "fake-unit-test-key")
    rejected = ORIGINAL.copy()
    rejected[8] = "Partial"
    manager = workbook(tmp_path, [rejected])
    gmail = Gmail()
    cmd_sync(manager, "FY 2026-27", gmail, leo_only=True, workers=1)
    assert gmail.fetched == []
    assert manager.wb.active.cell(2, 9).value == "Partial"


def test_completed_pdf_payloads_are_released_before_checkpoint(tmp_path, monkeypatch):
    """A long repair must not retain every processed PDF through its Futures."""
    import gc
    import weakref

    class Documents(dict):
        pass

    monkeypatch.setenv("OPENROUTER_API_KEY", "fake-unit-test-key")
    manager = workbook(tmp_path, [ORIGINAL.copy() for _ in range(25)])
    payloads = []

    def worker(row, sb, missing, gmail, cache=None, refresh=False, metrics=None):
        docs = Documents(finalLeo=("sb22.pdf", b"pdf-payload"))
        payloads.append(weakref.ref(docs))
        return row, sb, missing, docs, extracted(), None

    original_save = manager.save
    released_at_checkpoint = []

    def save():
        if not released_at_checkpoint:
            gc.collect()
            released_at_checkpoint.append(sum(ref() is None for ref in payloads))
        original_save()

    monkeypatch.setattr("sync._fetch_and_extract_worker", worker)
    monkeypatch.setattr(manager, "save", save)
    cmd_sync(manager, "FY 2026-27", Gmail(), force=True, leo_only=True, workers=1)
    assert released_at_checkpoint[0] >= 18
