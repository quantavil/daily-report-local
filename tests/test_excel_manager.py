"""Smoke tests for Excel hardening. No creds, no network needed."""

from datetime import date
from pathlib import Path

import openpyxl

from excel_manager import ExcelManager, excel_cell_to_str

HEADERS = [
    "Shipping Bill No", "Inv No", "LEO Date", "Shipper Name",
    "Container No.", "Container Size", "Customs Seal Number",
]


def test_cell_to_str_kills_dot_zero_poisoning():
    assert excel_cell_to_str(1278254.0) == "1278254"
    assert excel_cell_to_str(1278254) == "1278254"
    assert excel_cell_to_str(" 1278254 ") == "1278254"
    assert excel_cell_to_str(None) == ""
    assert excel_cell_to_str(date(2025, 4, 28)) == "28-Apr-2025"
    assert excel_cell_to_str(12.5) == "12.5"


def _workbook(path, sb_value):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FY 2026-27"
    ws.append(HEADERS)
    ws.append([sb_value, None, None, None, None, None, None])
    wb.save(path)
    return path


def test_float_sb_number_still_parses(tmp_path):
    xlsx = _workbook(tmp_path / "t.xlsx", 1278254.0)
    mgr = ExcelManager(xlsx)
    rows, _ = mgr.find_missing_rows("FY 2026-27")
    assert [r.sb_no for r in rows] == ["1278254"]


def test_save_creates_timestamped_backup(tmp_path):
    xlsx = _workbook(tmp_path / "t.xlsx", 1278254)
    mgr = ExcelManager(xlsx)
    mgr.load()
    mgr.save()
    backups = sorted(tmp_path.glob("t.bak-*.xlsx"))
    assert len(backups) >= 1


def test_backup_prunes_to_newest_five(tmp_path):
    xlsx = _workbook(tmp_path / "t.xlsx", 1278254)
    for i in range(7):
        (tmp_path / f"t.bak-2020010{i}-000000.xlsx").write_bytes(b"x")
    mgr = ExcelManager(xlsx)
    mgr.load()
    mgr.save()
    backups = sorted(tmp_path.glob("t.bak-*.xlsx"))
    assert len(backups) == 5
    assert not (tmp_path / "t.bak-20200100-000000.xlsx").exists()


def _workbook_with_remarks(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FY 2026-27"
    ws.append([*HEADERS, "Additional Remarks"])
    ws.append(["SB1", None, None, None, None, None, None, "keep me"])
    wb.save(path)
    return path


def test_ensure_status_column_replaces_additional_remarks(tmp_path):
    xlsx = _workbook_with_remarks(tmp_path / "t.xlsx")
    mgr = ExcelManager(xlsx)
    mgr.load()
    idx = mgr.ensure_status_column("FY 2026-27")
    ws = mgr.wb["FY 2026-27"]
    assert ws.cell(row=1, column=idx).value == "Status"
    assert idx == len(HEADERS) + 1  # same slot Additional Remarks had
    hdr = ws.cell(row=1, column=idx)
    assert hdr.fill.fgColor.rgb == "FF276738"  # distinct from other headers
    assert hdr.font.bold is True
    assert hdr.font.color.rgb == "FFFFFFFF"
    assert mgr.write_status("FY 2026-27", 2, idx, "Success")
    assert ws.cell(row=2, column=idx).value == "Success"
    # second call reuses the column instead of adding another
    assert mgr.ensure_status_column("FY 2026-27") == idx


def test_create_template_sheet_idempotent(tmp_path):
    xlsx = _workbook(tmp_path / "t.xlsx", 1278254)
    mgr = ExcelManager(xlsx)
    mgr.load()
    assert mgr.create_template_sheet(source="FY 2026-27") is True
    assert mgr.create_template_sheet(source="FY 2026-27") is False
    assert mgr.wb.sheetnames[0] == "Template"
    t = mgr.wb["Template"]
    assert t.max_row == 1  # headers only, no data rows
    assert t.cell(row=1, column=1).value == "Shipping Bill No"


def test_coerce_excel_date_formats():
    from excel_manager import coerce_excel_date

    # Standard DD-Mon-YYYY
    assert coerce_excel_date("28-Apr-2025") == date(2025, 4, 28)
    # Single digit day
    assert coerce_excel_date("6-Apr-2025") == date(2025, 4, 6)
    # Two digit year
    assert coerce_excel_date("28-Apr-25") == date(2025, 4, 28)
    # Slash format DD/MM/YYYY
    assert coerce_excel_date("28/04/2025") == date(2025, 4, 28)
    assert coerce_excel_date("6/4/2025") == date(2025, 4, 6)
    # ISO format YYYY-MM-DD
    assert coerce_excel_date("2025-04-28") == date(2025, 4, 28)
    # Unmatched format passes through
    assert coerce_excel_date("invalid-date") == "invalid-date"
    assert coerce_excel_date(None) is None


def test_apply_updates_preserves_formulas(tmp_path):
    from models import CustomsExtractedFields

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FY 2026-27"
    ws.append(HEADERS)
    # Add a row where Inv No is a formula
    ws.append(["1278254", "=VLOOKUP(A2, Sheet2!A:B, 2, FALSE)", None, None, None, None, None])
    path = tmp_path / "formula.xlsx"
    wb.save(path)

    mgr = ExcelManager(path)
    mgr.load()
    col_map = mgr.resolve_columns(mgr.wb["FY 2026-27"])

    updates = CustomsExtractedFields(
        invoice_numbers="OVERWRITE-ATTEMPT",
        shipper_name="NEW SHIPPER",
        leo_date="28-Apr-2025",
    )
    applied = mgr.apply_updates("FY 2026-27", 2, updates, col_map, force_overwrite=True)
    target_ws = mgr.wb["FY 2026-27"]

    # Formula was preserved
    assert target_ws.cell(row=2, column=col_map["inv_no"]).value == "=VLOOKUP(A2, Sheet2!A:B, 2, FALSE)"
    # Other fields applied
    assert target_ws.cell(row=2, column=col_map["shipper_name"]).value == "NEW SHIPPER"
    assert target_ws.cell(row=2, column=col_map["leo_date"]).value == date(2025, 4, 28)
    assert "inv_no" not in [a.split("=")[0] for a in applied]


def test_merged_cell_handling(tmp_path):
    from excel_manager import build_merged_cell_map, get_effective_cell_value
    from models import CustomsExtractedFields

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FY 2026-27"
    ws.append(HEADERS)
    ws.append(["1111111", "INV1", "28-Apr-2025", "SHIPPER 1", "CONT1234567", "40", "SEAL1"])
    ws.append(["2222222", "INV2", "28-Apr-2025", "SHIPPER 2", None, None, None])
    # Merge container columns between rows 2 and 3 (simulating LCL shipment)
    ws.merge_cells("E2:E3")  # Container No.
    ws.merge_cells("F2:F3")  # Container Size
    ws.merge_cells("G2:G3")  # Customs Seal
    path = tmp_path / "merged.xlsx"
    wb.save(path)

    mgr = ExcelManager(path)
    mgr.load()
    ws_loaded = mgr.wb["FY 2026-27"]
    m_map = build_merged_cell_map(ws_loaded)

    # Row 3 container cells should resolve to row 2's values
    assert get_effective_cell_value(ws_loaded, 3, 5, m_map) == "CONT1234567"
    assert get_effective_cell_value(ws_loaded, 3, 6, m_map) == "40"
    assert get_effective_cell_value(ws_loaded, 3, 7, m_map) == "SEAL1"

    # find_missing_rows should NOT flag row 3 as missing container fields
    missing, col_map = mgr.find_missing_rows("FY 2026-27")
    assert len(missing) == 0

    # apply_updates on row 3 should not crash with MergedCell read-only AttributeError
    updates = CustomsExtractedFields(
        container_numbers="CONT1234567",
        container_sizes="40",
        customs_seal_numbers="SEAL1",
    )
    applied = mgr.apply_updates("FY 2026-27", 3, updates, col_map, force_overwrite=False)
    assert len(applied) == 0  # skipped because already effectively non-empty


def test_is_cancelled_value():
    from excel_manager import is_cancelled_value

    assert is_cancelled_value("CANCELLED") is True
    assert is_cancelled_value("cancelled") is True
    assert is_cancelled_value("CANCELED") is True
    assert is_cancelled_value("CANCEL") is True
    assert is_cancelled_value("BILL CANCELLED") is True
    assert is_cancelled_value(None) is False
    assert is_cancelled_value("") is False
    assert is_cancelled_value("186") is False
    assert is_cancelled_value("PAID") is False


def test_find_missing_rows_skips_cancelled(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FY 2026-27"
    ws.append([*HEADERS, "Billing Status"])
    # Row 2: active with missing fields -> should be in missing_rows
    ws.append(["1278254", None, None, None, None, None, None, "186"])
    # Row 3: cancelled with missing fields -> should be SKIPPED
    ws.append(["2082481", None, None, None, None, None, None, "CANCELLED"])
    path = tmp_path / "cancelled.xlsx"
    wb.save(path)

    mgr = ExcelManager(path)
    mgr.load()
    missing, _ = mgr.find_missing_rows("FY 2026-27")
    assert len(missing) == 1
    assert missing[0].sb_no == "1278254"


def test_is_empty_cell_case_insensitive():
    from excel_manager import is_empty_cell

    # Truthy empties (including lowercase variants)
    assert is_empty_cell(None) is True
    assert is_empty_cell("") is True
    assert is_empty_cell("   ") is True
    assert is_empty_cell("-") is True
    assert is_empty_cell(" - ") is True
    assert is_empty_cell("NA") is True
    assert is_empty_cell("na") is True
    assert is_empty_cell("N/A") is True
    assert is_empty_cell("n/a") is True
    assert is_empty_cell("NONE") is True
    assert is_empty_cell("none") is True
    assert is_empty_cell("None") is True

    # Real non-empty values
    assert is_empty_cell("ACME EXPORTS") is False
    assert is_empty_cell("1278254") is False
    assert is_empty_cell(123) is False
    assert is_empty_cell(date(2026, 9, 10)) is False


def test_backup_microsecond_precision(tmp_path):
    xlsx = _workbook(tmp_path / "micro.xlsx", 1278254)
    mgr = ExcelManager(xlsx)
    mgr.load()
    # Rapid successive saves should generate distinct backup timestamps
    mgr.save()
    mgr.save()
    mgr.save()
    backups = list(tmp_path.glob("micro.bak-*.xlsx"))
    assert len(backups) == 3


def test_find_missing_rows_coerces_string_sb_date(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FY 2026-27"
    ws.append(["Shipping Bill No", "Inv No", "LEO Date", "Shipper Name",
               "Container No.", "Container Size", "Customs Seal Number", "Shipping Bill Date"])
    # Row 2 has sb_date as string "09-May-2024"
    ws.append(["1278254", None, None, None, None, None, None, "09-May-2024"])
    path = tmp_path / "string_sb_date.xlsx"
    wb.save(path)

    mgr = ExcelManager(path)
    mgr.load()
    missing, _ = mgr.find_missing_rows("FY 2026-27")
    assert len(missing) == 1
    assert missing[0].sb_date == date(2024, 5, 9)


def test_coerce_excel_date_extended_formats():
    from excel_manager import coerce_excel_date

    assert coerce_excel_date("28-04-2025") == date(2025, 4, 28)
    assert coerce_excel_date("28.04.2025") == date(2025, 4, 28)
    assert coerce_excel_date("6.4.2025") == date(2025, 4, 6)
    assert coerce_excel_date("28.Apr.2025") == date(2025, 4, 28)
    assert coerce_excel_date("2025/04/28") == date(2025, 4, 28)


def test_is_empty_cell_additional_placeholders():
    from excel_manager import is_empty_cell

    assert is_empty_cell("NULL") is True
    assert is_empty_cell("null") is True
    assert is_empty_cell("NIL") is True
    assert is_empty_cell("nil") is True
    assert is_empty_cell("--") is True
    assert is_empty_cell("NOT APPLICABLE") is True


def test_resolve_columns_with_linebreaks():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FY 2026-27"
    ws.append(["Shipping Bill\nNo", "Inv\nNo", "LEO\nDate", "Shipper\nName"])
    mgr = ExcelManager(Path("dummy.xlsx"))
    mgr.wb = wb
    col_map = mgr.resolve_columns(ws)
    assert col_map["sb_no"] == 1
    assert col_map["inv_no"] == 2
    assert col_map["leo_date"] == 3
    assert col_map["shipper_name"] == 4


def test_apply_updates_rejects_string_leo_date(tmp_path):
    from models import CustomsExtractedFields

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FY 2026-27"
    ws.append(HEADERS)
    ws.append(["1278254", None, None, None, None, None, None])
    path = tmp_path / "reject_leo.xlsx"
    wb.save(path)

    mgr = ExcelManager(path)
    mgr.load()
    col_map = mgr.resolve_columns(mgr.wb["FY 2026-27"])

    updates = CustomsExtractedFields(
        shipper_name="ACME EXPORTS",
        leo_date="Pending LEO",  # String that fails date coercion
    )
    applied = mgr.apply_updates("FY 2026-27", 2, updates, col_map, force_overwrite=True)
    target_ws = mgr.wb["FY 2026-27"]
    assert target_ws.cell(row=2, column=col_map["leo_date"]).value is None
    assert "leo_date" not in [a.split("=")[0] for a in applied]


def test_apply_updates_sets_wrap_text_on_multiline_cells(tmp_path):
    from models import CustomsExtractedFields

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FY 2026-27"
    ws.append(HEADERS)
    ws.append(["1278254", None, None, None, None, None, None])
    path = tmp_path / "wrap.xlsx"
    wb.save(path)

    mgr = ExcelManager(path)
    mgr.load()
    col_map = mgr.resolve_columns(mgr.wb["FY 2026-27"])

    updates = CustomsExtractedFields(
        container_numbers="TCKU7809338\nUNSU0088111",
    )
    mgr.apply_updates("FY 2026-27", 2, updates, col_map)
    cell = mgr.wb["FY 2026-27"].cell(row=2, column=col_map["container_no"])
    assert cell.alignment.wrap_text is True


def test_write_status_handles_merged_cells(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FY 2026-27"
    ws.append([*HEADERS, "Status"])
    ws.append(["1111111", None, None, None, None, None, None, None])
    ws.append(["2222222", None, None, None, None, None, None, None])
    ws.merge_cells("H2:H3")
    path = tmp_path / "merged_status.xlsx"
    wb.save(path)

    mgr = ExcelManager(path)
    mgr.load()
    status_col = len(HEADERS) + 1
    # Writing to merged child row 3 should write to top-left cell row 2 without crashing
    assert mgr.write_status("FY 2026-27", 3, status_col, "Success") is True
    assert mgr.wb["FY 2026-27"].cell(row=2, column=status_col).value == "Success"


def test_create_template_sheet_fallback_source(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DefaultSheet"
    ws.append(HEADERS)
    path = tmp_path / "tmpl_fallback.xlsx"
    wb.save(path)

    mgr = ExcelManager(path)
    mgr.load()
    # Source sheet 'NonExistent' falls back to active sheet without throwing KeyError
    assert mgr.create_template_sheet(name="Template", source="NonExistent") is True
    assert "Template" in mgr.wb.sheetnames

