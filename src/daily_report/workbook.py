"""Read workbook identities and apply scoped, sequential, validated updates."""
from __future__ import annotations

import os
import shutil
import tempfile
from copy import copy
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from daily_report import rules
from daily_report.models import BillIdentity, RowUpdate, WorkbookRow


def _header(value) -> str:
    return ' '.join(str(value or '').split()).casefold().rstrip('.:')


def _blank(value) -> bool:
    return value is None or str(value).strip().upper() in rules.EXCEL_BLANK_VALUES


def _date(value) -> date | None:
    if _blank(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in rules.EXCEL_DATE_FORMATS:
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    raise ValueError(f'Invalid date: {value!r}')


def _sb(value) -> str | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, float):
        if not value.is_integer():
            return None
        value = int(value)
    text = str(value).strip()
    if text.endswith('.0'):
        text = text[:-2]
    if not text.isascii():
        return None
    try:
        return rules.value_token('sb_no', text)
    except ValueError:
        return None


class WorkbookStore:
    """An in-memory workbook; reads never cause persistence or header changes."""

    def __init__(self, path: Path):
        self.path = Path(path)
        self.wb = openpyxl.load_workbook(self.path, data_only=False, read_only=False)
        self.dirty = False

    def _sheet(self, sheet: str) -> Worksheet:
        if sheet not in self.wb.sheetnames:
            raise ValueError(f'Unknown sheet: {sheet}')
        return self.wb[sheet]

    @staticmethod
    def _columns(ws: Worksheet) -> dict[str, int]:
        headers = [_header(cell.value) for cell in ws[1]]
        columns = {}
        for field, aliases in rules.EXCEL_COLUMNS.items():
            matches = [i + 1 for i, header in enumerate(headers) if header in {_header(a) for a in aliases}]
            if len(matches) > 1:
                raise ValueError(f'Ambiguous columns for {field} in {ws.title}')
            if matches:
                columns[field] = matches[0]
        missing = set(rules.EXCEL_REQUIRED_COLUMNS) - columns.keys()
        if missing:
            raise ValueError(f'Missing required columns in {ws.title}: {", ".join(sorted(missing))}')
        return columns

    @staticmethod
    def _merged(ws: Worksheet) -> tuple[dict, dict]:
        parents = {}
        covered_rows = {}
        for region in ws.merged_cells.ranges:
            parent = region.min_row, region.min_col
            covered_rows[parent] = set(range(region.min_row, region.max_row + 1))
            for row in range(region.min_row, region.max_row + 1):
                for col in range(region.min_col, region.max_col + 1):
                    parents[row, col] = parent
        return parents, covered_rows

    def rows(self, sheet: str, row_numbers: set[int] | None = None,
             sb_no: str | None = None) -> list[WorkbookRow]:
        ws = self._sheet(sheet)
        columns = self._columns(ws)
        parents, _ = self._merged(ws)
        rows = []
        for row_idx in range(2, ws.max_row + 1):
            if row_numbers is not None and row_idx not in row_numbers:
                continue

            def value(field, row_idx=row_idx):
                if field not in columns:
                    return None
                position = parents.get((row_idx, columns[field]), (row_idx, columns[field]))
                return ws.cell(*position).value

            sb = _sb(value('sb_no'))
            if sb is None or (sb_no is not None and sb != sb_no):
                continue
            if str(value('billing_status') or '').strip().upper() == rules.EXCEL_CANCELLED_VALUE:
                rows.append(WorkbookRow(row_idx=row_idx, identity=BillIdentity(sb_no=sb),
                                        values={}, cancelled=True))
                continue
            try:
                identity = BillIdentity(sb_no=sb, sb_date=_date(value('sb_date')),
                                        port_code=None if _blank(value('port_code')) else value('port_code'),
                                        document_date=_date(value('document_date')))
            except ValueError as exc:
                rows.append(WorkbookRow(row_idx=row_idx, identity=BillIdentity(sb_no=sb),
                                        values={}, error=f'{sheet} row {row_idx}: {exc}'))
                continue
            values = {}
            for field in rules.DATA_FIELDS:
                raw = value(field)
                if _blank(raw):
                    values[field] = None
                elif field in rules.EXCEL_DATE_FIELDS and not str(raw).startswith('='):
                    try:
                        values[field] = _date(raw)
                    except ValueError:
                        values[field] = str(raw)
                else:
                    values[field] = str(raw) if not isinstance(raw, date) else raw
            rows.append(WorkbookRow(row_idx=row_idx, identity=identity, values=values))
        return rows

    def apply(self, sheet: str, updates: list[RowUpdate], force: bool = False,
              fields: set[str] | None = None) -> dict[int, list[str]]:
        ws = self._sheet(sheet)
        columns = self._columns(ws)
        if not updates:
            return {}
        selected = {update.row_idx for update in updates}
        if len(selected) != len(updates):
            raise ValueError('Duplicate row updates')
        if any(row > ws.max_row for row in selected):
            raise ValueError('Update row is outside the workbook data range')
        allowed = set(rules.DATA_FIELDS) if fields is None else set(fields)
        if allowed - set(rules.DATA_FIELDS):
            raise ValueError('Unknown writable fields: ' + ', '.join(sorted(allowed - set(rules.DATA_FIELDS))))
        parents, covered_rows = self._merged(ws)
        cancelled_rows = set()
        if 'billing_status' in columns:
            for row in range(2, ws.max_row + 1):
                position = parents.get((row, columns['billing_status']), (row, columns['billing_status']))
                if str(ws.cell(*position).value or '').strip().upper() == rules.EXCEL_CANCELLED_VALUE:
                    cancelled_rows.add(row)
        writable_rows = selected - cancelled_rows
        problems: dict[int, list[str]] = {}
        by_row = {update.row_idx: update for update in updates}
        plans: dict[tuple[int, int], list[tuple[int, str, object]]] = {}

        def flag(row, message):
            if message not in problems.setdefault(row, []):
                problems[row].append(message)
            update = by_row[row]
            if update.status not in {'Cancelled', 'Error'}:
                update.status = 'Partial'
            update.detail = '; '.join(filter(None, [update.detail, message]))

        for update in updates:
            if update.row_idx in cancelled_rows:
                update.status = 'Cancelled'
                continue
            for field, value in update.values.items():
                if field not in rules.DATA_FIELDS:
                    raise ValueError(f'Unknown writable field: {field}')
                if isinstance(value, list) and field not in rules.CONTAINER_ATTRS:
                    raise ValueError('List values are only supported for container fields')
                if field not in allowed or _blank(value):
                    continue
                if update.containers is not None and field in rules.CONTAINER_ATTRS:
                    value = '\n'.join(getattr(record, rules.CONTAINER_ATTRS[field]) or ''
                                      for record in update.containers)
                    if _blank(value):
                        continue
                elif isinstance(value, list):
                    value = '\n'.join(item or '' for item in value)
                    if _blank(value):
                        continue
                if field in rules.EXCEL_DATE_FIELDS:
                    value = _date(value)
                target = parents.get((update.row_idx, columns[field]), (update.row_idx, columns[field]))
                plans.setdefault(target, []).append((update.row_idx, field, value))

        blocked = set()
        for target, proposed in plans.items():
            involved = {row for row, _, _ in proposed}
            if target[0] not in writable_rows or covered_rows.get(target, {target[0]}) - writable_rows:
                blocked.add(target)
                touches_cancelled = bool(covered_rows.get(target, {target[0]}) & cancelled_rows)
                reason = 'cancelled' if touches_cancelled else 'unselected'
                for row in involved:
                    flag(row, f'Merged cell {ws.cell(*target).coordinate} touches {reason} rows')
            if len({value for _, _, value in proposed}) > 1:
                blocked.add(target)
                for row in involved:
                    flag(row, f'Conflicting updates for {ws.cell(*target).coordinate}')

        # Status receives the same scope and conflict protection as data cells.
        if 'status' not in columns:
            columns['status'] = ws.max_column + 1
            header = ws.cell(1, columns['status'])
            header.value = rules.EXCEL_COLUMNS['status'][0]
            if header.column > 1:
                header._style = copy(ws.cell(1, header.column - 1)._style)
            self.dirty = True
        statuses = {}
        for update in updates:
            target = parents.get((update.row_idx, columns['status']), (update.row_idx, columns['status']))
            statuses.setdefault(target, []).append(update)
        blocked_status = set()
        for target, candidates in statuses.items():
            if target[0] not in selected or covered_rows.get(target, {target[0]}) - selected:
                blocked_status.add(target)
                for update in candidates:
                    flag(update.row_idx, f'Merged status {ws.cell(*target).coordinate} touches unselected rows')
            elif len({update.status for update in candidates}) > 1:
                blocked_status.add(target)
                for update in candidates:
                    flag(update.row_idx, f'Conflicting updates for status {ws.cell(*target).coordinate}')

        for target, proposed in plans.items():
            if target in blocked:
                continue
            cell = ws.cell(*target)
            if cell.data_type == 'f' or (not force and not _blank(cell.value)):
                continue
            value = proposed[0][2]
            if cell.value == value:
                continue
            cell.value = value
            if isinstance(value, date):
                cell.number_format = rules.EXCEL_DATE_NUMBER_FORMAT
            if isinstance(value, str) and '\n' in value:
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                cell.alignment = alignment
            self.dirty = True
        for target, candidates in statuses.items():
            if target in blocked_status:
                continue
            cell = ws.cell(*target)
            if cell.data_type != 'f' and cell.value != candidates[0].status:
                cell.value = candidates[0].status
                self.dirty = True
        return problems

    def save(self) -> None:
        if not self.dirty:
            return
        fd, temporary = tempfile.mkstemp(prefix=f'.{self.path.stem}-', suffix='.tmp.xlsx', dir=self.path.parent)
        os.close(fd)
        temporary = Path(temporary)
        try:
            self.wb.save(temporary)
            with temporary.open('rb') as handle:
                os.fsync(handle.fileno())
            stamp = datetime.now().strftime('%Y%m%d-%H%M%S-%f')
            backup = self.path.with_name(f'{self.path.stem}.bak-{stamp}{self.path.suffix}')
            shutil.copy2(self.path, backup)
            os.chmod(backup, 0o600)
            os.chmod(temporary, self.path.stat().st_mode & 0o777)
            os.replace(temporary, self.path)
            self.dirty = False
            backups = sorted(self.path.parent.glob(f'{self.path.stem}.bak-*{self.path.suffix}'))
            for stale in backups[:-5]:
                stale.unlink()
        finally:
            temporary.unlink(missing_ok=True)
